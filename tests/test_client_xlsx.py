import io
import unittest

from openpyxl import load_workbook

from ual_app.client_xlsx import client_xlsx_bytes, clean_csv


SAMPLE = b"Name,Empty,Country,Country\nAlice,,Canada,Ontario\n\nBob,,United States,Texas\nFormula,,=2+2,\n"


class ClientXlsxTests(unittest.TestCase):
    def test_clean_csv_removes_empty_rows_and_columns(self):
        headers, rows = clean_csv(SAMPLE)
        self.assertEqual(headers, ["Name", "Country", "Country (2)"])
        self.assertEqual(len(rows), 3)

    def test_workbook_has_client_table_and_hidden_unused_area(self):
        data, filename = client_xlsx_bytes("evidence.csv", SAMPLE)
        workbook = load_workbook(io.BytesIO(data), data_only=False)
        sheet = workbook["Client Data"]
        table = sheet.tables["ClientShareableData"]
        self.assertEqual(filename, "evidence-client-shareable.xlsx")
        self.assertEqual(table.ref, "A1:C4")
        self.assertEqual(table.tableStyleInfo.name, "TableStyleMedium2")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertTrue(sheet.sheet_format.zeroHeight)
        self.assertTrue(sheet.column_dimensions["D"].hidden)
        self.assertEqual(sheet["B4"].value, "=2+2")
        self.assertEqual(sheet["B4"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
